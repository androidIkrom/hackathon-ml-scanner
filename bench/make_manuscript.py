"""Builds the full manuscript (Word) and its figures from the measured results.

    python bench/make_manuscript.py [data-dir] [file-name]    (default D:/nungil-bench, Nungil-manuscript.docx)

Run analyze.py first. Every measured number is read from <data-dir>/results.json. References are numbered in
the order they are first cited. Text between [[ and ]] is a citation key; {{...}} marks a blank the authors fill in.
"""
import json
import re
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from matplotlib.patches import Arc, FancyBboxPatch, Wedge

import make_article as base
from make_article import BLUE, GRID, INK, MUTED, NAMES, ORANGE, ORDER, PUBLISHED, SECONDARY, pct, save, style_axes

REPO = Path(__file__).resolve().parents[1]
ACCESSED = "accessed 3 October 2026"
FIGURES = ["architecture", "direction", "evaluation", "latency", "tradeoff", "sizes", "threshold", "classes", "scores", "ablation", "gallery", "facethreshold", "screens"]
TABLES = ["research", "apps", "components", "environment", "detection", "facepipe", "ablation", "gallery", "tuning", "published", "metadata"]


def F(key):
    return f"Figure {FIGURES.index(key) + 1}"


def T(key):
    return f"Table {TABLES.index(key) + 1}"


REFS = {
    "who2026": f"World Health Organization. Blindness and vision impairment, https://www.who.int/news-room/fact-sheets/detail/blindness-and-visual-impairment (2026, {ACCESSED}).",
    "bourne2021": "GBD 2019 Blindness and Vision Impairment Collaborators. Trends in prevalence of blindness and distance and near vision impairment over 30 years: an analysis for the Global Burden of Disease Study. Lancet Glob Health 2021; 9: e130–e143.",
    "khan2021": "Khan A and Khusro S. An insight into smartphone-based assistive solutions for visually impaired and blind people: issues, challenges and opportunities. Univ Access Inf Soc 2021; 20: 265–298.",
    "kuriakose2022": "Kuriakose B, Shrestha R and Sandnes FE. Tools and technologies for blind and visually impaired navigation support: a review. IETE Tech Rev 2022; 39: 3–18.",
    "budrionis2022": "Budrionis A, Plikynas D, Daniušis P, et al. Smartphone-based computer vision travelling aids for blind and visually impaired individuals: a systematic review. Assist Technol 2022; 34: 178–194.",
    "manship": "Ngnamsie Njimbouom S, Lee K and Kim JD. MANSHIP: mobile-based assistive notification service for hearing-impaired people using a hybrid deep learning model. Technol Health Care 2025; 33: 1787–1799.",
    "lin2017": "Lin BS, Lee CC and Chiang PY. Simple smartphone-based guiding system for visually impaired people. Sensors 2017; 17: 1371.",
    "bai2019": "Bai J, Liu Z, Lin Y, et al. Wearable travel aid for environment perception and navigation of visually impaired people. Electronics 2019; 8: 697.",
    "joshi2020": "Joshi RC, Yadav S, Dutta MK, et al. Efficient multi-object detection and smart navigation using artificial intelligence for visually impaired people. Entropy 2020; 22: 941.",
    "see2022": "See AR, Sasing BG and Advincula WD. A smartphone-based mobility assistant using depth imaging for visually impaired and blind. Appl Sci 2022; 12: 2802.",
    "kuriakose2023": "Kuriakose B, Shrestha R and Sandnes FE. DeepNAVI: a deep learning based smartphone navigation assistant for people with visual impairments. Expert Syst Appl 2023; 212: 118720.",
    "zhang2024": "Zhang K, Wang Y, Shi S, et al. Improved yolov5 algorithm combined with depth camera and embedded system for blind indoor visual assistance. Sci Rep 2024; 14: 23000.",
    "chen2024": "Chen Z, Liu Z, Wang K, et al. A large vision-language model based environment perception system for visually impaired people. In: 2024 IEEE/RSJ international conference on intelligent robots and systems (IROS), 2024, pp.221–228. IEEE.",
    "gao2025": "Gao Y, Wu D, Song J, et al. A wearable obstacle avoidance device for visually impaired individuals with cross-modal learning. Nat Commun 2025; 16: 2857.",
    "mekonnen2025": "Mekonnen FY, Al Bataineh MF, Abu Abdoun D, et al. Development of a fully autonomous offline assistive system for visually impaired individuals: a privacy-first approach. Sensors 2025; 25: 6006.",
    "walkvlm": "Yuan Z, Zhang T, Zhu Y, et al. WalkVLM: aid visually impaired people walking by vision language model. In: 2025 IEEE/CVF international conference on computer vision (ICCV), 2025, pp.9845–9854. IEEE.",
    "alhammad2026": "Alhammad N, Alsubaie A, Alomair R, et al. Munir: a multimodal smart-glasses system for enhancing human–computer interaction for visually impaired individuals. Sensors 2026; 26: 3950.",
    "lin2026": "Lin HY, Fan YH and Chang CC. Multimodal navigation system for visually impaired users using environmental perception and vision-language models. Sensors 2026; 26: 3045.",
    "visionaid": "Florea CG and Spînu S. VisionAId: an offline-first multimodal Android assistant for people with visual impairment, featuring personalized object retrieval. arXiv preprint arXiv:2607.02371, 2026.",
    "seeingai_store": f"Microsoft. Seeing AI (App Store listing), https://apps.apple.com/us/app/seeing-ai/id999062298 ({ACCESSED}).",
    "seeingai_blog": f"Microsoft. Seeing AI app launches on Android, https://blogs.microsoft.com/accessibility/seeing-ai-app-launches-on-android-including-new-and-updated-features-and-new-languages/ (2023, {ACCESSED}).",
    "seeingai_manual": f"Microsoft. Seeing AI user manual, revision D, https://www.microsoft.com/en-us/garage/wp-content/uploads/2025/04/Seeing-AI-User-Manual.pdf (2024, {ACCESSED}).",
    "lookout_help": f"Google. Lookout help, https://support.google.com/accessibility/android/answer/9031274 ({ACCESSED}).",
    "lookout_blog": f"Google. Lookout announcement, https://blog.google/company-news/outreach-and-initiatives/accessibility/lookout-app-help-blind-and-visually-impaired-people-learn-about-their-surroundings/ (2018, {ACCESSED}).",
    "bme_store": f"Be My Eyes. Be My Eyes (App Store listing), https://apps.apple.com/us/app/be-my-eyes/id905177575 ({ACCESSED}).",
    "bme_site": f"Be My Eyes, https://www.bemyeyes.com/ ({ACCESSED}).",
    "bme_ai": f"Be My Eyes. Using Be My AI, https://support.bemyeyes.com/hc/en-us/articles/18133134809105-Using-Be-My-AI ({ACCESSED}).",
    "bme_meta": f"Meta. Be My Eyes on AI glasses, https://www.meta.com/help/ai-glasses/1524316621830425/ ({ACCESSED}).",
    "envision_app": f"Envision. Envision app, https://www.letsenvision.com/app ({ACCESSED}).",
    "envision_store": f"Envision. Envision AI (App Store listing), https://apps.apple.com/us/app/envision-ai/id1268632314 ({ACCESSED}).",
    "envision_android": f"Envision. Android release notes, https://www.letsenvision.com/updates/android ({ACCESSED}).",
    "envision_instant": f"Envision. Instant Text, https://www.letsenvision.com/blog/instant-text (2023, {ACCESSED}).",
    "sullivan_us": f"TUAT. Sullivan+ (App Store listing, United States), https://apps.apple.com/us/app/sullivan/id1475001234 ({ACCESSED}).",
    "sullivan_kr": f"TUAT. Sullivan+ (App Store listing, Republic of Korea), https://apps.apple.com/kr/app/%EC%84%A4%EB%A6%AC%EB%B2%88/id1475001234 ({ACCESSED}).",
    "sullivan_play": f"TUAT. Sullivan+ (Google Play listing), https://play.google.com/store/apps/details?id=tuat.kr.sullivan ({ACCESSED}).",
    "sullivan_blog": f"TUAT. Sullivan+ functions, https://blog.mysullivan.org/2019/09/sullivan-plus-about-functions.html (2019, {ACCESSED}).",
    "tts_site": f"TapTapSee, https://taptapseeapp.com/ ({ACCESSED}).",
    "tts_instr": f"TapTapSee. Instructions, https://taptapseeapp.com/instruction.html ({ACCESSED}).",
    "tts_store": f"TapTapSee (App Store listing), https://apps.apple.com/us/app/taptapsee/id567635020 ({ACCESSED}).",
    "mediapipe": "Lugaresi C, Tang J, Nash H, et al. MediaPipe: a framework for building perception pipelines. arXiv preprint arXiv:1906.08172, 2019.",
    "efficientdet": "Tan M, Pang R and Le QV. EfficientDet: scalable and efficient object detection. In: 2020 IEEE/CVF conference on computer vision and pattern recognition (CVPR), 2020, pp.10778–10787. IEEE.",
    "ssd": "Liu W, Anguelov D, Erhan D, et al. SSD: single shot multibox detector. In: Computer vision – ECCV 2016. Cham: Springer, 2016, pp.21–37.",
    "mobilenetv2": "Sandler M, Howard A, Zhu M, et al. MobileNetV2: inverted residuals and linear bottlenecks. In: 2018 IEEE/CVF conference on computer vision and pattern recognition, 2018, pp.4510–4520. IEEE.",
    "efficientnet": "Tan M and Le QV. EfficientNet: rethinking model scaling for convolutional neural networks. arXiv preprint arXiv:1905.11946, 2019.",
    "mlkit": f"Google. ML Kit, https://developers.google.com/ml-kit ({ACCESSED}).",
    "facenet": "Schroff F, Kalenichenko D and Philbin J. FaceNet: a unified embedding for face recognition and clustering. In: 2015 IEEE conference on computer vision and pattern recognition (CVPR), 2015, pp.815–823. IEEE.",
    "facenet_android": f"shubham0204. FaceRecognition_With_FaceNet_Android (GitHub repository), https://github.com/shubham0204/FaceRecognition_With_FaceNet_Android ({ACCESSED}).",
    "mobilenetv3": "Howard A, Sandler M, Chen B, et al. Searching for MobileNetV3. In: 2019 IEEE/CVF international conference on computer vision (ICCV), 2019, pp.1314–1324. IEEE.",
    "arcoredepth": "Valentin J, Kowdle A, Barron JT, et al. Depth from motion for smartphone AR. ACM Trans Graph 2018; 37: 1–19.",
    "coco": "Lin TY, Maire M, Belongie S, et al. Microsoft COCO: common objects in context. In: Computer vision – ECCV 2014. Cham: Springer, 2014, pp.740–755.",
    "lfw": "Huang GB, Ramesh M, Berg T, et al. Labeled faces in the wild: a database for studying face recognition in unconstrained environments. Technical Report 07-49, University of Massachusetts, Amherst, 2007.",
    "efron1979": "Efron B. Bootstrap methods: another look at the jackknife. Ann Stat 1979; 7: 1–26.",
    "wilson1927": "Wilson EB. Probable inference, the law of succession, and statistical inference. J Am Stat Assoc 1927; 22: 209–212.",
    "tfblog": f"TensorFlow Blog. Easier object detection on mobile with TensorFlow Lite, https://blog.tensorflow.org/2021/06/easier-object-detection-on-mobile-with-tf-lite.html (2021, {ACCESSED}).",
    "mp_objdet": f"Google. Object detection task guide (MediaPipe Solutions), https://developers.google.com/edge/mediapipe/solutions/vision/object_detector ({ACCESSED}).",
    "deepface": f"Serengil SI. DeepFace benchmarks, https://github.com/serengil/deepface/tree/master/benchmarks ({ACCESSED}).",
}


class Cites:
    def __init__(self):
        self.order = []

    def number(self, key):
        if key not in REFS:
            raise KeyError(f"unknown reference {key}")
        if key not in self.order:
            self.order.append(key)
        return self.order.index(key) + 1

    def label(self, keys):
        numbers = sorted({self.number(k.strip()) for k in keys.split(",")})
        parts, start, last = [], numbers[0], numbers[0]
        for n in numbers[1:] + [None]:
            if n is not None and n == last + 1:
                last = n
                continue
            parts.append(str(start) if start == last else f"{start}–{last}" if last > start + 1 else f"{start},{last}")
            start = last = n
        return ",".join(parts)


CITES = Cites()


# ---------- figures ----------

def box(ax, x, y, w, h, text, fill, size=7.4, bold=False):
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0,rounding_size=1.2", linewidth=0.8, edgecolor="#898781", facecolor=fill))
    ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=size, color=INK, fontweight="bold" if bold else "normal", linespacing=1.25)


def fig_architecture(path):
    fig, ax = plt.subplots(figsize=(6.6, 5.0))
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 100)
    ax.axis("off")
    rows = [
        ("Sensors", 82, 12, "#f1f0ec", ["Camera\n(CameraX / ARCore)", "Rotation-vector\nsensor (heading)", "Depth map\n(ARCore)", "Microphone", "GPS"]),
        ("On-device models", 56, 15, "#cde2fb", ["Object detector:\nEfficientDet-Lite,\nSSD MobileNet V2", "Face detector\n+ FaceNet-512", "Image embedder:\nMobileNetV3-small", "Text and barcode\nrecognition", "Speech\nrecogniser"]),
        ("Core logic (pure Kotlin, unit-tested)", 29, 15, "#ffffff", ["Direction,\n360°\ncoverage", "Duplicate\nmerging,\ncounting", "Colour\nnaming", "Face and\nobject\nmatching", "Find\nguidance", "Walk\nhazards", "Voice\ncommands"]),
        ("Output", 4, 12, "#f1f0ec", ["Speech (English / Korean)", "Beeps", "Vibration", "On-screen captions"]),
    ]
    for title, y, h, fill, cells in rows:
        ax.text(1, y + h + 2.2, title, fontsize=8.2, color=SECONDARY, fontweight="bold", va="bottom")
        gap = 1.6
        w = (98 - gap * (len(cells) - 1)) / len(cells)
        for i, cell in enumerate(cells):
            box(ax, 1 + i * (w + gap), y, w, h, cell, fill)
    for upper, lower in ((82, 56 + 15 + 6.2), (56, 29 + 15 + 6.2), (29, 4 + 12 + 6.2)):
        for x in (20, 50, 80):
            ax.annotate("", xy=(x, lower), xytext=(x, upper - 0.4), arrowprops=dict(arrowstyle="-|>", color="#52514e", linewidth=1.1))
    return save(fig, path)


def fig_direction(path):
    fig, ax = plt.subplots(figsize=(5.2, 3.6))
    ax.set_aspect("equal")
    ax.set_xlim(-0.55, 1.75)
    ax.set_ylim(-0.12, 1.32)
    ax.axis("off")
    heading, fov, offset = 35.0, 65.0, 18.0        # degrees, clockwise from the start of the scan

    def point(compass, radius):
        a = np.radians(90 - compass)
        return radius * np.cos(a), radius * np.sin(a)

    ax.add_patch(Wedge((0, 0), 1.15, 90 - (heading + fov / 2), 90 - (heading - fov / 2), facecolor="#cde2fb", edgecolor="#86b6ef", linewidth=0.8))
    ax.plot([0, 0], [0, 1.2], color=MUTED, linewidth=1, linestyle=(0, (4, 3)))
    ax.text(0, 1.24, "start of scan (0°)", ha="center", fontsize=8, color=SECONDARY)
    hx, hy = point(heading, 1.2)
    ax.annotate("", xy=(hx, hy), xytext=(0, 0), arrowprops=dict(arrowstyle="-|>", color=INK, linewidth=1.3))
    ax.text(hx + 0.03, hy + 0.03, "camera axis", fontsize=8, color=INK)
    ox, oy = point(heading + offset, 0.92)
    ax.plot([0, ox], [0, oy], color=ORANGE, linewidth=1.2, linestyle=(0, (2, 2)))
    ax.plot([ox], [oy], marker="o", markersize=8, color=ORANGE, markeredgecolor="white", markeredgewidth=1.2)
    ax.text(ox + 0.06, oy - 0.02, "object", fontsize=8.5, color=INK, va="center")
    ax.add_patch(Arc((0, 0), 0.62, 0.62, theta1=90 - heading, theta2=90, color=INK, linewidth=1))
    ax.text(*point(heading / 2, 0.38), "h", fontsize=9.5, color=INK, ha="center", va="center", style="italic")
    ax.add_patch(Arc((0, 0), 1.08, 1.08, theta1=90 - heading - offset, theta2=90 - heading, color=ORANGE, linewidth=1.2))
    ax.text(*point(heading + offset + 19, 0.5), "(x − 0.5)·FOV", fontsize=8, color=INK, ha="left", va="center")
    ex, ey = point(heading + fov / 2, 1.0)
    ax.text(ex + 0.05, ey - 0.03, "edge of the field\nof view (FOV)", fontsize=8, color=SECONDARY, va="top")
    ax.plot([0], [0], marker="s", markersize=9, color=INK)
    ax.text(-0.06, -0.02, "phone", fontsize=8.5, color=INK, ha="right", va="center")
    ax.text(0.98, 0.1, "θ = h + (x − 0.5)·FOV", fontsize=9.5, color=INK, style="italic")
    return save(fig, path)


def fig_evaluation(path, images, faces):
    fig, ax = plt.subplots(figsize=(6.6, 1.75))
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 30)
    ax.axis("off")
    cells = [
        ("Public datasets", f"COCO val2017 sample\n({images:,} images)\nLFW ({faces:,} images)", "#f1f0ec"),
        ("Phone", "instrumented\nbenchmark runs the\napp's detector and\nface pipeline", "#cde2fb"),
        ("Raw output", "detections and times,\nface embeddings", "#f1f0ec"),
        ("Analysis", "AP, precision, recall,\nverification, open-set\nidentification, 95% CI", "#ffffff"),
    ]
    gap, w = 5.0, (98 - 5.0 * 3) / 4
    for i, (title, text, fill) in enumerate(cells):
        x = 1 + i * (w + gap)
        box(ax, x, 2, w, 21, text, fill, size=7.2)
        ax.text(x + w / 2, 25, title, ha="center", fontsize=8.2, color=SECONDARY, fontweight="bold")
        if i:
            ax.annotate("", xy=(x - 0.3, 12.5), xytext=(x - gap + 0.3, 12.5), arrowprops=dict(arrowstyle="-|>", color="#52514e", linewidth=1.1))
    return save(fig, path)


def fig_threshold(path, sweep):
    fig, ax = plt.subplots(figsize=(6.3, 3.0))
    xs = [row["threshold"] for row in sweep]
    for key, colour, name in (("precision", BLUE, "Precision"), ("recall", ORANGE, "Recall")):
        ys = [100 * row[key] for row in sweep]
        ax.plot(xs, ys, color=colour, linewidth=2, marker="o", markersize=5, label=name)
        for x, y in zip(xs, ys):
            ax.annotate(f"{y:.1f}", (x, y), xytext=(0, 7), textcoords="offset points", ha="center", fontsize=8.5, color=INK)
    ax.axvline(0.7, color=MUTED, linewidth=1, linestyle=(0, (4, 3)))
    ax.text(0.695, 6, "deployed default", ha="right", fontsize=8.5, color=SECONDARY)
    ax.set_xlabel("Announcement threshold on the detector score (people 0.2 lower)")
    ax.set_ylabel("%")
    ax.set_xticks(xs)
    ax.set_xlim(0.27, 0.73)
    ax.set_ylim(0, 108)
    style_axes(ax, "y")
    ax.legend(frameon=False, fontsize=9, loc="center", bbox_to_anchor=(0.6, 0.56))
    return save(fig, path)


def fig_classes(path, rows):
    fig, ax = plt.subplots(figsize=(6.3, 5.0))
    for index, row in enumerate(rows):
        for offset, key, colour in ((-0.19, "precision", BLUE), (0.19, "recall", ORANGE)):
            if row[key] is None:      # nothing of this class was announced, so precision is undefined
                ax.text(1, index + offset, "none announced", va="center", fontsize=7.8, color=SECONDARY, style="italic")
                continue
            value = 100 * row[key]
            ax.barh(index + offset, value, height=0.34, color=colour, label={"precision": "Precision", "recall": "Recall"}[key] if index == 0 else None)
            ax.text(value + 1, index + offset, f"{value:.0f}", va="center", fontsize=7.8, color=INK)
    ax.set_yticks(range(len(rows)), [f"{row['label']} (n = {row['objects']})" for row in rows])
    ax.invert_yaxis()
    ax.set_xlabel("% at the application's thresholds")
    ax.set_xlim(0, 112)
    ax.set_xticks(range(0, 101, 20))
    style_axes(ax, "x")
    ax.legend(frameon=False, fontsize=9, ncol=2, loc="lower center", bbox_to_anchor=(0.5, 1.0))
    return save(fig, path)


def fig_scores(path, scores):
    fig, ax = plt.subplots(figsize=(6.3, 2.9))
    edges = scores["edges"]
    centres = [(a + b) / 2 for a, b in zip(edges, edges[1:])]
    width = edges[1] - edges[0]
    ax.bar(centres, scores["different"], width=width * 0.92, color=ORANGE, alpha=0.85, label="Different people")
    ax.bar(centres, scores["same"], width=width * 0.92, color=BLUE, alpha=0.75, label="Same person")
    ax.axvline(0.5, color=INK, linewidth=1, linestyle=(0, (4, 3)))
    ax.text(0.485, max(scores["same"]) * 0.96, "threshold 0.5", fontsize=8.5, color=INK, ha="right")
    ax.set_xlabel("Cosine similarity of the two face embeddings")
    ax.set_ylabel("Number of pairs")
    style_axes(ax, "y")
    ax.legend(frameon=False, fontsize=9, loc="upper left")
    return save(fig, path)


def fig_ablation(path, rows):
    short = ["Single sample ≥ 0.3", "Single sample ≥ 0.5", "Mean of 3 best ≥ 0.5", "+ margin 0.08 (deployed)"]
    fig, axes = plt.subplots(1, 2, figsize=(6.6, 2.5), sharey=True)
    for ax, key, title, limits in ((axes[0], "named_right", "Saved person named correctly (%)", (98, 100.3)),
                                   (axes[1], "stranger_named", "Stranger given a saved name (%)", (0, 98))):
        for index, row in enumerate(rows):
            value, (low, high) = 100 * row[key], row[key + "_ci"]
            ax.errorbar([value], [index], xerr=[[value - 100 * low], [100 * high - value]], color=BLUE, marker="o", markersize=6, capsize=3, ecolor=SECONDARY, elinewidth=1)
            ax.annotate(f"{value:.1f}", (value, index), xytext=(0, 8), textcoords="offset points", ha="center", fontsize=8.5, color=INK)
        ax.set_title(title, fontsize=9.5, color=INK)
        ax.set_xlim(*limits)
        ax.set_ylim(len(rows) - 0.5, -0.7)
        style_axes(ax, "x")
    axes[0].set_yticks(range(len(rows)), short)
    return save(fig, path)


def fig_facethreshold(path, rows):
    fig, ax = plt.subplots(figsize=(6.3, 3.0))
    xs = [row["threshold"] for row in rows]
    for key, colour, name in (("named_right", BLUE, "Saved person named correctly"), ("stranger_named", ORANGE, "Stranger given a saved name")):
        ys = [100 * row[key] for row in rows]
        ax.plot(xs, ys, color=colour, linewidth=2, marker="o", markersize=5, label=name)
        for x, y in zip(xs, ys):
            ax.annotate(f"{y:.1f}", (x, y), xytext=(0, 7), textcoords="offset points", ha="center", fontsize=8.5, color=INK)
    ax.axvline(0.5, color=MUTED, linewidth=1, linestyle=(0, (4, 3)))
    ax.text(0.505, 45, "deployed", fontsize=8.5, color=SECONDARY)
    ax.set_xlabel("Threshold on the person score (10 saved people, 5 samples each)")
    ax.set_ylabel("%")
    ax.set_xticks(xs)
    ax.set_xlim(0.37, 0.73)
    ax.set_ylim(0, 110)
    style_axes(ax, "y")
    ax.legend(frameon=False, fontsize=9, loc="center right")
    return save(fig, path)


def figures(results, folder):
    plt.rcParams.update({"font.family": "DejaVu Sans", "text.color": INK, "axes.labelcolor": SECONDARY})
    paths = base.figures(results, folder)
    extras = results["extras"]
    images = results["detection"][0]["images"]
    return {
        "architecture": fig_architecture(folder / "m-architecture.png"), "direction": fig_direction(folder / "m-direction.png"),
        "evaluation": fig_evaluation(folder / "m-evaluation.png", images, results["faces"]["images"]),
        "latency": paths["latency"], "tradeoff": paths["tradeoff"], "sizes": paths["sizes"], "gallery": paths["faces"],
        "threshold": fig_threshold(folder / "m-threshold.png", extras["detection_threshold_sweep"]),
        "classes": fig_classes(folder / "m-classes.png", extras["detection_per_class"]),
        "scores": fig_scores(folder / "m-scores.png", extras["face_pair_scores"]),
        "ablation": fig_ablation(folder / "m-ablation.png", extras["face_rule_ablation"]),
        "facethreshold": fig_facethreshold(folder / "m-facethreshold.png", extras["face_threshold"]),
    }


import numpy as np  # noqa: E402  (used by fig_direction)


# ---------- Word helpers ----------

WORDS = [0]


def rich(paragraph, text, size=None, bold=None, italic=None, count=True):
    """Adds text to a paragraph: [[keys]] become superscript reference numbers, <i> and <b> switch style, {{...}} is a blank."""
    state = {"i": bool(italic), "b": bool(bold)}
    for part in re.split(r"(\[\[.*?\]\]|</?[ib]>|\{\{.*?\}\})", text):
        if not part:
            continue
        if part.startswith("[["):
            run = paragraph.add_run(CITES.label(part[2:-2]))
            run.font.superscript = True
        elif part in ("<i>", "</i>", "<b>", "</b>"):
            state[part.strip("</>")] = not part.startswith("</")
            continue
        else:
            run = paragraph.add_run(part)
            if part.startswith("{{"):
                run.font.highlight_color = 7   # yellow: a blank for the authors
            elif count:
                WORDS[0] += len(part.split())
        run.bold = state["b"] or None
        run.italic = state["i"] or None
        if size:
            run.font.size = Pt(size)
    return paragraph


class Doc:
    def __init__(self):
        self.doc = Document()
        normal = self.doc.styles["Normal"]
        normal.font.name, normal.font.size = "Times New Roman", Pt(12)
        normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Malgun Gothic")
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.5
        for name, size in (("Title", 16), ("Heading 1", 13), ("Heading 2", 12), ("Heading 3", 12)):
            style = self.doc.styles[name]
            style.font.name, style.font.size, style.font.bold = "Times New Roman", Pt(size), True
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.font.italic = name == "Heading 3"
            style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:asciiTheme"), "")
            for attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                style.element.rPr.rFonts.attrib.pop(qn(attribute), None)
            style.element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
            style.element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        border = self.doc.styles["Title"].element.pPr.find(qn("w:pBdr"))
        if border is not None:
            self.doc.styles["Title"].element.pPr.remove(border)
        section = self.doc.sections[0]
        section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Cm(2.5)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)

    def heading(self, text, level=1):
        WORDS[0] += 0
        return self.doc.add_heading(text, level=level)

    def para(self, text, **kwargs):
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        return rich(paragraph, text, **kwargs)

    def bullets(self, items, style="List Bullet", **kwargs):
        for item in items:
            rich(self.doc.add_paragraph(style=style), item, **kwargs)

    def caption(self, label, text):
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.line_spacing = 1.0
        rich(paragraph, f"<b>{label}.</b> {text}", size=10)

    def table(self, header, rows, widths, size=9, first_bold=False, note=None):
        grid = self.doc.add_table(rows=1, cols=len(header))
        grid.style = "Table Grid"
        grid.autofit = False
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        grid.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row_index, values in enumerate([header] + rows):
            cells = grid.rows[0].cells if row_index == 0 else grid.add_row().cells
            for column, (cell, value) in enumerate(zip(cells, values)):
                cell.width = Cm(widths[column])
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                rich(paragraph, str(value), size=size, bold=row_index == 0 or (first_bold and column == 0), count=False)
                if row_index == 0:
                    base.shade(cell, "EDEDEA")
        if note:
            paragraph = self.doc.add_paragraph()
            paragraph.paragraph_format.line_spacing = 1.0
            rich(paragraph, note, size=9, count=False)
        else:
            self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def figure(self, path, key, text, width=15.5):
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.add_run().add_picture(str(path), width=Cm(width))
        caption = self.doc.add_paragraph()
        caption.paragraph_format.line_spacing = 1.0
        rich(caption, f"<b>{F(key)}.</b> {text}", size=10)

    def algorithm(self, title, lines):
        grid = self.doc.add_table(rows=2, cols=1)
        grid.style = "Table Grid"
        head, body = grid.rows[0].cells[0], grid.rows[1].cells[0]
        base.shade(head, "EDEDEA")
        head.paragraphs[0].paragraph_format.space_after = Pt(0)
        head.paragraphs[0].paragraph_format.line_spacing = 1.0
        rich(head.paragraphs[0], f"<b>{title}</b>", size=10, count=False)
        for index, line in enumerate(lines):
            paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(line)
            run.font.name, run.font.size = "Consolas", Pt(8.5)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def equation(self, text, number):
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rich(paragraph, f"<i>{text}</i>        ({number})", count=False)

    def orient(self, landscape):
        section = self.doc.add_section(WD_SECTION.NEW_PAGE)
        wide, narrow = max(section.page_width, section.page_height), min(section.page_width, section.page_height)
        section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = (wide, narrow) if landscape else (narrow, wide)

    def page_break(self):
        self.doc.add_page_break()


def versions():
    gradle = (REPO / "app" / "build.gradle").read_text(encoding="utf-8") + (REPO / "build.gradle").read_text(encoding="utf-8")
    find = lambda pattern: (re.search(pattern, gradle) or [None, "?"])[1]
    return {
        "mediapipe": find(r"mediapipe:tasks-vision:([\d.]+)"), "face": find(r"mlkit:face-detection:([\d.]+)"), "tflite": find(r"tensorflow-lite:([\d.]+)"),
        "text": find(r"mlkit:text-recognition:([\d.]+)"), "barcode": find(r"mlkit:barcode-scanning:([\d.]+)"), "camerax": find(r"camera-core:([\d.]+)"),
        "arcore": find(r"ar:core:([\d.]+)"), "room": find(r"room-runtime:([\d.]+)"), "kotlin": find(r"kotlin\.android' version '([\d.]+)"),
        "agp": find(r"com\.android\.application' version '([\d.]+)"), "min": find(r"minSdk (\d+)"), "target": find(r"targetSdk (\d+)"),
    }


RESEARCH_HEADER = ["System", "Platform", "Inference", "Functions", "Model(s)", "Evaluation data", "Reported accuracy", "Reported latency (device)", "User study"]
RESEARCH = [
    ["Lin et al., 2017[[lin2017]]", "Android phone + PC server with GPU", "Mixed: server online, phone offline mode", "OD, Dir, Dist, Obs", "Faster R-CNN, YOLO",
     "Own campus images + Pascal, 7 classes", "Recognition rate 55% (7 classes), 60% (5 classes)", "Qualitative only", "4 visually impaired (partial sight)"],
    ["Bai et al., 2019[[bai2019]]", "RGB-D + IMU glasses + smartphone (Snapdragon 820)", "On-device", "Obs, walkable direction, OD, Nav",
     "MobileNetV2 detector trained on COCO; depth ground segmentation", "Office field test", "Not reported", "114 ms object detection (phone)", "20 visually impaired"],
    ["Joshi et al., 2020[[joshi2020]]", "Single-board processor (quad-core 1.5 GHz), camera, ultrasonic sensor", "On-device", "OD with count, Obs, banknotes", "YOLOv3",
     "Own dataset, 150 test images per class", "Detection accuracy 95.19%, recognition accuracy 99.69%", "0.3 s per frame without GPU, 0.1 s with GPU", "36 (20 visually impaired, 16 blindfolded)"],
    ["See et al., 2022[[see2022]] (abstract only)", "Smartphone with depth camera", "Not accessible", "Obs, OD, Find, Voice and gestures", "Not accessible", "Not accessible",
     "Objects located \"more than 80% of the time\"", "Not accessible", "Not accessible"],
    ["DeepNAVI, 2023[[kuriakose2023]]", "Android smartphone + bone-conduction headset", "On-device", "OD (20 obstacle types), Dir, Dist, Scene",
     "EfficientDet-Lite4, EfficientNet-Lite4", "Own 20-class obstacle and scene datasets", "Obstacle detection accuracy 87.8%, scene 85%", "Not reported", "Pilot with 1 visually impaired user"],
    ["Zhang et al., 2024[[zhang2024]]", "Raspberry Pi 4B + RealSense D435i", "On-device", "Find (object named by voice), Dist, Voice", "Improved YOLOv5",
     "Own 1,835 images, 3 classes", "Precision 91.7%, recall 94.2%, mAP 93.9%", "4.3 ms (hardware not stated)", "Not reported"],
    ["Chen et al., 2024[[chen2024]]", "Wearable touch device + cloud server (A100 GPU)", "Cloud", "Scene, object category by touch", "Qwen-VL-Chat + SETR segmentation",
     "POPE, MME, LLaVA-QA90", "POPE accuracy 0.86–0.89", "0.4–4.4 s per action, end to end", "12 visually impaired"],
    ["Gao et al., 2025[[gao2025]]", "Custom glasses (FPGA, video + depth) + smartphone", "Mixed: compression on glasses, detection on phone", "Obs (audio + tactile)",
     "YOLOv5-based cross-modal detector", "720 indoor and outdoor trials", "100% collision avoidance rate", "< 320 ms end to end", "20 with impaired vision"],
    ["Mekonnen et al., 2025[[mekonnen2025]]", "Raspberry Pi 5", "On-device", "OD, OCR, Face (enrolment), Voice", "YOLOv8n, Tesseract, face_recognition",
     "Own controlled tests", "OD precision 85.2%, recall 82.7%; face 81.3%; OCR words 89.7%", "OD 0.80 s per frame; face 0.89 s; OCR 1.17 s", "None"],
    ["WalkVLM, 2025[[walkvlm]]", "Vision-language model on cloud devices", "Cloud", "Walking reminders, scene questions", "MiniCPM-V2.6 + LoRA",
     "Walking Awareness Dataset, 1.5k test samples", "Text-similarity scores (TF-IDF 0.166)", "Not reported", "34 surveys; closed-loop test with 3 visually impaired people"],
    ["Munir, 2026[[alhammad2026]]", "ESP32-S3 smart glasses + phone app (Helio G91-Ultra)", "Mixed: face on-device; other functions in the cloud",
     "Face (enrolment), OCR, Scene, OD, Voice", "InsightFace buffalo_l; cloud vision and language APIs", "SCface (1,690 probes) + 30 LFW impostor probes",
     "Face recognition rate 96.69%, false accept rate 0%", "Face 44.3 ms on phone; 3.3 s end to end", "None"],
    ["Lin et al., 2026[[lin2026]]", "Head-mounted RealSense D455 + laptop", "Mixed: local perception, remote language model", "Nav (indoor), OD (open vocabulary), Obs, Scene, Voice",
     "RTAB-Map, YOLO-World, GPT-4o-mini", "RePOPE; indoor trials", "RePOPE accuracy 87.9–90.9%", "3.3 s mean guidance response", "None"],
]
APPS_HEADER = ["Feature", "Nungil (this work)", "Seeing AI[[seeingai_store,seeingai_blog,seeingai_manual]]", "Lookout[[lookout_help,lookout_blog]]",
               "Be My Eyes[[bme_store,bme_site,bme_ai,bme_meta]]", "Envision[[envision_app,envision_store,envision_android,envision_instant]]",
               "Sullivan+[[sullivan_us,sullivan_kr,sullivan_play,sullivan_blog]]", "TapTapSee[[tts_site,tts_instr,tts_store]]"]
APPS = [
    ["Platform", "Android 7.0+", "iOS, Android", "Android 6+", "iOS, Android", "iOS, Android", "iOS, Android", "iOS, Android"],
    ["Where processing runs", "On the device; only the optional walking-route request uses the network", "Not stated",
     "\"Core experience\" on device (2018); food labels offline", "Be My AI requires internet", "Instant Text offline; other modes not stated", "Not stated", "Cloud: image is uploaded"],
    ["Live object detection", "Yes, 80 COCO classes", "iOS with LiDAR only", "Explore mode (beta)", "—", "Explore", "AI mode; live operation unverified", "No; photo or short video"],
    ["Direction of objects", "Compass heading + camera field of view, spoken", "Spatial audio", "Direction and distance in Find mode", "—", "Yes", "Spoken location", "—"],
    ["360° room summary", "Yes", "—", "—", "—", "—", "—", "—"],
    ["Colour naming", "Yes, 11 colour names", "Yes", "—", "—", "Yes", "Yes", "—"],
    ["Find a named object", "Beeps + vibration; object classes, saved people, saved objects", "Find My Things", "Find mode, 7 fixed categories", "—", "Find Objects, in-app list", "Object Finder by voice", "—"],
    ["Recognise saved faces", "Yes", "Yes", "—", "—", "Yes", "Unverified", "—"],
    ["Teach personal objects", "Yes (appearance embeddings)", "Yes; Android unverified", "—", "—", "—", "—", "—"],
    ["Text, QR, barcode", "Latin text, QR, barcode; no Korean text", "Text, barcode, QR", "Text, documents, food barcodes", "No dedicated mode", "Text, barcode; offline Korean text", "Text, PDF, QR", "Barcode, QR"],
    ["Walking hazard warnings", "Walls, steps, stairs, drop-offs (depth)", "—", "—", "—", "—", "—", "—"],
    ["Navigation to places", "Saved places; online route or offline direction beacon", "Indoor recorded routes (iOS)", "—", "—", "—", "—", "—"],
    ["Hands-free voice control", "Wake word; all functions", "Unverified", "Assistant launch; voice questions in English", "On Meta glasses", "Voice Mode in Ask Envision", "Unverified", "—"],
    ["Published accuracy or latency", f"This work", "None", "None", "None found", "Relative only (\"30% faster\")", "None found", "\"7 to 10 seconds\" per photo"],
]


def build(results, paths, target):
    det = {(e["model"], e["delegate"]): e for e in results["detection"]}
    faces, extras = results["faces"], results["extras"]
    ver, ident = faces["verification"], faces["identification"]["by_saved_people"]
    L2, L2C = det[("efficientdet-lite2.tflite", "GPU")], det[("efficientdet-lite2.tflite", "CPU")]
    L0, I8, SSD = det[("efficientdet-lite0.tflite", "GPU")], det[("efficientdet-lite0-int8.tflite", "CPU")], det[("ssd-mobilenet-v2.tflite", "GPU")]
    precisions = [e["precision"] for e in det.values()]
    rule, tune, enrol, sweep, classes, scores = (extras[k] for k in ("face_rule_ablation", "face_threshold", "face_enrolment", "detection_threshold_sweep", "detection_per_class", "face_pair_scores"))
    objects = L2["tp"] + L2["fn"]
    v = versions()
    person = next(row for row in classes if row["label"] == "person")
    chair = next(row for row in classes if row["label"] == "chair")
    highlights = [
        "All recognition runs on the phone: no cloud service and no extra hardware",
        f"Deployed detector: {100 * L2['precision']:.0f}% precision at {100 * L2['recall']:.0f}% recall on COCO, {L2['ms_mean']:.0f} ms per image",
        f"A three-part matching rule cuts strangers named as saved people from {100 * rule[0]['stranger_named']:.0f}% to {100 * rule[3]['stranger_named']:.0f}%",
        "An open benchmark runs the app's own pipeline on COCO and LFW on the phone",
        "Results carry 95% confidence intervals and the deployed settings, not model scores",
    ]
    for line in highlights:
        assert len(line) <= 85, (len(line), line)

    d = Doc()
    d.doc.add_heading("Nungil: A fully on-device smartphone scene assistant for blind and low-vision people, evaluated on public benchmarks at its deployed settings", level=0)
    d.para("{{Author names, in order}}", count=False)
    d.para("{{Affiliations: department, institution, city, country}}", count=False)
    d.para("<b>Corresponding author:</b> {{name, postal address, email}}", count=False)

    d.heading("Abstract")
    d.para("<b>Background:</b> Smartphone assistants for blind and low-vision people usually depend on a cloud service or on extra hardware, and published systems typically report accuracy on their own data, which makes them hard to compare.")
    d.para("<b>Objective:</b> To present Nungil, an Android application that performs all recognition on the phone, and to measure its deployed pipeline on public benchmarks.")
    d.para("<b>Methods:</b> Nungil combines object detection with compass-based direction, colour naming and a de-duplicated summary of a full turn; find-object guidance; recognition of saved people and personal objects; text reading; and depth-based hazard warnings, under voice control in English and Korean. "
           f"An instrumented benchmark ran the application's own detector and face pipeline on a mid-range phone over {L2['images']:,} COCO val2017 images and {faces['images']:,} Labeled Faces in the Wild images.")
    d.para(f"<b>Results:</b> At the deployed thresholds the default detector reached a precision of {pct(L2['precision'])} (95% CI {100 * L2['precision_ci'][0]:.1f}–{100 * L2['precision_ci'][1]:.1f}) and a recall of {pct(L2['recall'])} overall and {pct(L2['recall_by_size']['large'])} on large objects, at {L2['ms_mean']:.0f} ms per image. "
           f"Face verification accuracy was {100 * ver['ten_fold_accuracy']:.2f}%. With ten saved people, a saved person was named correctly in {pct(ident['10']['named_right'])} of cases and {pct(ident['10']['stranger_named'])} of strangers received a saved name; "
           f"under a single-sample rule that share was {pct(rule[0]['stranger_named'])}.")
    d.para("<b>Conclusions:</b> A phone alone can give announcements that are almost always right, at the price of missing small objects. Reporting deployed settings on public benchmarks with confidence intervals makes assistive systems comparable.")
    d.para("<b>Keywords:</b> assistive technology, visual impairment, on-device inference, object detection, face recognition, smartphone", count=False)

    # ---------------- 1 Introduction ----------------
    d.heading("Introduction")
    d.para("At least 2.2 billion people worldwide have a near or distance vision impairment.[[who2026]] In 2020 an estimated 43.3 million people were blind and 295 million had moderate or severe vision impairment.[[bourne2021]] "
           "For these people, everyday tasks that sighted people solve at a glance need help: knowing what is in a room, finding a bag, recognising who is standing in front of them, reading a sign, or noticing a step before reaching it.")
    d.para("The smartphone is an attractive platform for such help because it already combines a camera, motion sensors, speech input and output, and enough computing power to run neural networks. "
           "Smartphone-based assistive solutions have been surveyed repeatedly,[[khan2021,kuriakose2022,budrionis2022]] and mobile assistive services are being built for other sensory impairments as well.[[manship]] "
           "Several commercial applications are widely used, among them Seeing AI,[[seeingai_store]] Lookout,[[lookout_help]] Be My Eyes[[bme_site]] and Envision.[[envision_app]]")
    d.para("Three gaps remain. First, many systems depend on a server, a cloud model or dedicated hardware such as smart glasses, depth cameras or single-board computers, which raises cost, needs connectivity and sends images of the user's surroundings off the device. "
           "Second, published systems are evaluated on their own datasets with their own metrics, so their results cannot be compared, and they usually report the accuracy of a model, not of what the application finally says with its thresholds and filters. "
           "Third, recognition of familiar people is an open-set problem: a stranger must not be greeted by a saved name, yet this error is rarely measured.")
    d.para("This paper presents Nungil (눈길), an Android application that addresses these gaps. Its contributions are:")
    d.bullets([
        "A fully on-device assistant that combines object detection with direction and colour, a spoken summary of a full 360° turn, find-object guidance, recognition of saved people and personal objects, text and code reading, and depth-based walking hazards, controlled by voice in English and Korean and released as open source.",
        "A sensor-fusion scan that gives every detected object a compass direction from the phone's heading and the camera's field of view, and a de-duplication rule that turns hundreds of frames into one correct count.",
        f"An open-set face matching rule (threshold, margin over the runner-up, mean of the three best samples) whose ablation shows that it lowers the share of strangers given a saved name from {pct(rule[0]['stranger_named'])} to {pct(rule[3]['stranger_named'])}.",
        "A reproducible benchmark that runs the application's own pipeline on the phone over COCO and Labeled Faces in the Wild and reports precision, recall, latency and identification rates at the deployed settings with 95% confidence intervals.",
    ])
    d.para("The rest of the paper reviews related work, describes the proposed system, reports the experiments, describes the implementation, discusses the results and their limitations, and concludes.")

    # ---------------- 2 Related work ----------------
    d.orient(landscape=True)
    d.heading("Related work")
    d.para(f"{T('research')} compares Nungil with twelve representative assistive systems for blind and visually impaired people published between 2017 and 2026. It is a selection, not a systematic review. "
           "Every cell for the other systems is taken from the cited paper; \"not reported\" means that the paper does not state it and \"not accessible\" that the text could not be read. "
           "OD = object detection, Dir = direction, Dist = distance, Find = find-object guidance, Face = face recognition, OCR = text reading, Obs = obstacle or hazard warning, Nav = navigation, Scene = scene description, Voice = voice control.")
    d.caption(T("research"), "Nungil and twelve published assistive systems: platform, functions and the results each paper reports.")
    ours = ["Nungil (this work)", "Android phone alone (Infinix X6880, Helio G99)", "On-device", "OD, Dir, colour, 360° summary, Find, Face (enrolment), personal objects, OCR, Obs, Nav, Voice",
            "EfficientDet-Lite0/Lite2, SSD MobileNet V2; FaceNet-512; MobileNetV3-small", f"COCO val2017 sample ({L2['images']:,} images); LFW ({faces['images']:,} images)",
            f"COCO AP {100 * L2['ap']:.1f}; precision {pct(L2['precision'])}, recall {pct(L2['recall'])} at deployed thresholds; LFW verification {100 * ver['ten_fold_accuracy']:.2f}%",
            f"{L2['ms_mean']:.0f} ms (Lite2, GPU); {I8['ms_mean']:.0f} ms (Lite0 int8, CPU); face {faces['find_ms_median']:.0f} + {faces['embed_ms_median']:.0f} ms", "None yet"]
    d.table(RESEARCH_HEADER, [ours] + RESEARCH, [2.8, 3.0, 1.9, 3.1, 2.7, 2.7, 3.4, 2.6, 2.4], size=7.5, first_bold=True)
    d.orient(landscape=False)

    d.heading("Research systems", level=2)
    d.para("The systems fall into three groups. Server- and cloud-based systems send images away for recognition: an early smartphone guide used a server with Faster R-CNN and YOLO,[[lin2017]] and recent ones use vision-language models to describe scenes, answer questions or give walking reminders.[[chen2024,walkvlm,lin2026]] "
           "They offer open-ended descriptions but need a connection, and where response times are reported they range from 0.4 to 4.4 s. "
           "Wearable and embedded systems add hardware: RGB-D glasses,[[bai2019]] single-board computers with ultrasonic or depth sensors,[[joshi2020,zhang2024,mekonnen2025]] FPGA glasses with a cross-modal detector,[[gao2025]] or smart glasses paired with a phone.[[alhammad2026]] "
           "Smartphone-only systems are the smallest group: a depth-camera phone for obstacle avoidance and object finding[[see2022]] and DeepNAVI, which runs EfficientDet-Lite4 on the device and reports obstacle type, direction and distance.[[kuriakose2023]]")
    d.para(f"Three observations follow from {T('research')}. First, none of the twelve papers reports results on a shared public benchmark such as COCO for detection or Labeled Faces in the Wild for faces; they use their own images, hallucination benchmarks or task success, so their figures cannot be ranked. "
           "Values such as a detection accuracy of 95.19%[[joshi2020]] or a mean average precision of 93.9%[[zhang2024]] were measured on a few custom classes. "
           "Second, none of the twelve reports colour naming, a summary of a full turn or user-taught personal objects, and face enrolment appears in only two.[[mekonnen2025,alhammad2026]] "
           "Third, several were tested with visually impaired participants,[[lin2017,bai2019,joshi2020,chen2024,gao2025]] which the present work has not yet done.")
    d.para("A concurrent preprint, VisionAId, is the closest in spirit: an offline-first Android assistant with six on-device models and a few-shot pipeline for personal objects, with an optional cloud language model for scene narration.[[visionaid]] It has not been peer reviewed; it reports a depth latency of about 491 ms and a banknote-detector mAP@50 of 0.986 on a Samsung Galaxy S21 Ultra, no public-benchmark results and no user study.")

    d.orient(landscape=True)
    d.heading("Commercial applications", level=2)
    d.para(f"{T('apps')} lists the documented features of six widely used applications. The cells come only from the official product pages, store listings, help pages and manuals cited in the column headers. "
           "A dash means that the feature is not listed on the pages reviewed, which does not prove that it is absent. None of the six publishes accuracy figures; one states a response time of 7 to 10 seconds per photo.")
    d.caption(T("apps"), "Documented features of Nungil and six assistive applications (public sources, retrieved October 2026).")
    d.table(APPS_HEADER, APPS, [2.6, 3.6, 3.0, 3.2, 2.6, 3.2, 3.2, 3.0], size=8, first_bold=True)
    d.orient(landscape=False)

    # ---------------- 3 Proposed system ----------------
    d.heading("Proposed system")
    d.heading("Overview", level=2)
    d.para(f"The user holds the phone upright and either turns around or walks; the application speaks what it finds. {F('architecture')} shows the architecture. "
           "Sensor data pass through neural models that run on the device and then through a layer of plain rules that decides what is worth saying. "
           "This rule layer is written in pure Kotlin without Android dependencies, so it can be unit-tested on a desktop machine. Speech, beeps, vibration and on-screen captions form the output. "
           f"{T('components')} lists the components and the operating points that were evaluated.")
    d.figure(paths["architecture"], "architecture", "Architecture of Nungil. Sensor data are processed by on-device models and by a rule layer in pure Kotlin; nothing is sent to a server.")
    d.caption(T("components"), "Components of the pipeline and their operating points.")
    d.table(["Function", "Method", "Operating point"], [
        ["Object detection", "MediaPipe object detector;[[mediapipe]] EfficientDet-Lite2 (default) and Lite0,[[efficientdet]] SSD[[ssd]] MobileNet V2;[[mobilenetv2]] EfficientDet-Lite0 int8 in walk mode",
         "Detector score ≥ 0.3, at most 10 results per frame; announced at ≥ 0.7 (people ≥ 0.5)"],
        ["Fallback naming", "EfficientNet-Lite0 image classifier[[efficientnet]] on the middle of the frame when no object is detected", "Score ≥ 0.35"],
        ["Direction", "Rotation-vector heading + box position × camera field of view", "4 sectors in speech, 8 internally"],
        ["De-duplication", "Clusters by label and angle; count = most seen in one frame", "Merge within 20°; confirmed after 3 frames"],
        ["Colour", "Gray-world white balance; per-pixel hue, saturation and value vote over 24 × 24 samples", "11 colour names; no colour in dark frames"],
        ["Saved people", "ML Kit face detection;[[mlkit]] levelled crop; FaceNet-512[[facenet,facenet_android]] on the crop and its mirror image", "Cosine ≥ 0.5, margin ≥ 0.08, mean of 3 best samples"],
        ["Saved objects", "MobileNetV3-small[[mobilenetv3]] image embeddings", "Cosine ≥ 0.75"],
        ["Reader", "ML Kit text recognition and barcode scanning", "Latin script, QR codes, barcodes"],
        ["Walk mode", "ARCore depth:[[arcoredepth]] walls, steps, stairs, drop-offs; detector for people and vehicles", "Readings between 0.3 and 8 m"],
    ], [3.0, 8.0, 5.0], size=8.5, first_bold=True)

    d.heading("Detection and direction", level=2)
    d.para("Each camera frame is passed to the object detector, which returns up to ten boxes with one of the 80 COCO classes. A box is announced only when its score reaches the user's setting (0.7 by default, adjustable down to 0.3); people are accepted 0.2 lower because distant and half-hidden people score low. "
           "A box that touches exactly one side edge of the frame is dropped, since the same object will be seen whole in another frame.")
    d.para(f"The direction of an object combines two measurements ({F('direction')}). The rotation-vector sensor gives the heading <i>h</i> of the camera axis relative to the start of the scan, and the horizontal position <i>x</i> of the box centre (0 at the left edge, 1 at the right) gives the offset inside the field of view:")
    d.equation("θ = (h + (x − 0.5) · FOV) mod 360°", 1)
    d.para("The horizontal field of view follows from the physical camera data, with <i>w</i> the sensor side that maps to the image width and <i>f</i> the focal length; 65° is used when the data are unusable:")
    d.equation("FOV = 2 · arctan(w / (2f))", 2)
    d.para("The angle is spoken as one of four sectors (front, right, behind, left). During a full scan the heading is also recorded in 36 bins of 10°; the scan ends when every bin has been covered or after 60 s, and the user is asked to slow down when turning faster than 60° per second.")
    d.figure(paths["direction"], "direction", "Direction of an object, seen from above: the heading of the camera axis plus the offset of the object inside the field of view.", width=11.5)

    d.heading("De-duplication and summary", level=2)
    d.para("During a turn the same chair is detected in dozens of frames. Algorithm 1 merges detections of the same class whose directions are within 20° and keeps, as the count of a cluster, the largest number seen in a single frame. "
           "Summing over frames would turn three chairs into three hundred. A cluster is spoken only after it has been seen in three frames, and its colour is the majority vote of its detections, so a flickering colour cannot split one object into two. "
           "The result is one sentence, for example \"Around you: 3 blue chairs in front; a black laptop on your right.\"")
    d.algorithm("Algorithm 1: De-duplication of detections during a scan", [
        "Input:  detections D of one frame (label, angle, colour); clusters C; merge angle τ = 20°; n = 3",
        "Output: clusters confirmed by this frame",
        " 1  seen ← empty map",
        " 2  for each detection d in D do",
        " 3      c ← the cluster in C with c.label = d.label and |c.angle − d.angle| ≤ τ, nearest in angle",
        " 4      if there is none then create c from d and add it to C",
        " 5      c.angle ← circular mean of c.angle (weight c.samples) and d.angle",
        " 6      c.samples ← c.samples + 1;  add d.colour to the votes of c",
        " 7      seen[c] ← seen[c] + 1",
        " 8  for each (c, k) in seen do",
        " 9      c.frames ← c.frames + 1",
        "10      c.count ← max(c.count, k)          ▷ most seen in one frame, never a sum",
        "11      if c.frames = n then report c       ▷ spoken only after n frames",
        "12  colour of c: the most voted colour, if it has ≥ 2 votes and more than any other",
    ])
    d.heading("Colour naming", level=2)
    d.para("The colour of an object is taken from 24 × 24 samples in the central half of its box. After a gray-world white balance each sample votes for one of eleven basic colour names according to its hue, saturation and value; "
           "a chromatic name wins once chromatic samples make up at least 30% of the named samples. No colour is given in dark frames, and people are never given a colour.")

    d.heading("Saved people", level=2)
    d.para("Faces are found with the ML Kit face detector. A face is used only when it is at least 64 pixels wide and turned no more than 35° sideways or 25° up or down. "
           "The crop is levelled by the line between the eyes and embedded twice with FaceNet-512, once as it is and once mirrored; the two embeddings are averaged and normalised. A person is enrolled in one guided sweep of the head (straight, left, right, up, down).")
    d.para("Naming a face is an open-set decision, because most people the camera sees are not saved. Algorithm 2 scores every saved person by the mean cosine similarity of the three samples closest to the probe, "
           "requires the best score to reach a threshold of 0.5, and requires the best person to be ahead of the runner-up by a margin of 0.08. Otherwise the face stays unnamed.")
    d.algorithm("Algorithm 2: Open-set matching of a face against the saved people", [
        "Input:  probe embedding v; gallery G = {person p → samples S_p}; threshold t = 0.5; margin m = 0.08; k = 3",
        "Output: a saved person, or none",
        " 1  for each person p in G do",
        " 2      score[p] ← mean of the k largest cos(v, s) for s in S_p",
        " 3  b ← person with the highest score;  r ← person with the second highest score",
        " 4  if score[b] < t then return none",
        " 5  if r exists and score[b] − score[r] < m then return none",
        " 6  return b",
    ])
    d.heading("Saved objects and find guidance", level=2)
    d.para("Personal objects are recognised by appearance: crops of the object are embedded with MobileNetV3-small and compared by cosine similarity with a threshold of 0.75. "
           "When the user asks to find something, by class, by a saved name or by a saved object, the phone beeps faster as the target moves towards the centre of the frame, from one beep per second at the edge to one every 150 ms at the centre, vibrates when it is centred, and says that it lost the target after 3 s without a match.")
    d.heading("Reader and walk mode", level=2)
    d.para("The reader speaks printed Latin text, QR codes and barcodes. In walk mode the depth map from ARCore replaces guesswork about distance. "
           "Readings between 0.3 and 8 m in the band from 30% to 62% of the image height are inspected in three zones; a zone counts as blocked when at least 40% of its readings are within 3 m, and as unknown when fewer than 20% of its readings are valid. "
           "A floor profile ahead of the user yields steps, stairs and drop-offs. The application speaks only when something changes and never says that the way is safe.")
    d.heading("Voice interaction", level=2)
    d.para("The application listens continuously and reacts after the wake word \"Eye\" (Korean \"눈길\"). All functions can be started by voice in English and Korean, everything spoken is also shown as a caption, and application sound stops while the user is speaking.")

    # ---------------- 4 Experimentation ----------------
    d.heading("Experimentation")
    d.heading("Datasets", level=2)
    d.para(f"Object detection was evaluated on a fixed random sample of {L2['images']:,} images of COCO val2017[[coco]] (seed 2026), which contain {objects:,} annotated objects other than crowd regions. "
           f"Face recognition was evaluated on all {faces['images']:,} images of Labeled Faces in the Wild (LFW),[[lfw]] with its standard list of 6,000 verification pairs. "
           "The detectors were trained by their publishers on COCO; to our knowledge the validation images used here were not part of that training, and neither dataset was used to choose the thresholds of the application.")
    d.heading("Environment setup", level=2)
    d.para(f"{T('environment')} gives the test environment. All measurements were taken on one mid-range phone connected to a charger.")
    d.caption(T("environment"), "Environment of the experiments.")
    d.table(["Item", "Value"], [
        ["Phone", "Infinix X6880; MediaTek MT6789 (Helio G99), Mali-G57 GPU; 8 GB RAM; Android 15"],
        ["Application", f"Nungil, commit b65c9b3 of the main branch plus local changes outside the evaluated code; Kotlin {v['kotlin']}, Android Gradle Plugin {v['agp']}, target SDK {v['target']}"],
        ["Inference libraries", f"MediaPipe Tasks Vision {v['mediapipe']}; ML Kit face detection {v['face']}; TensorFlow Lite {v['tflite']} (4 threads for face embedding)"],
        ["Models", "EfficientDet-Lite2 (448 × 448), EfficientDet-Lite0 (320 × 320), EfficientDet-Lite0 int8, SSD MobileNet V2, FaceNet-512"],
        ["Phone state", "Charging over USB; benchmark in the foreground; battery temperature 36–37 °C"],
        ["Analysis", "Python 3.13, NumPy, pycocotools 2.0.11; laptop with Intel Core i5-11400H"],
    ], [3.5, 12.5], size=9, first_bold=True)
    d.heading("Procedure and metrics", level=2)
    d.para(f"The evaluation follows {F('evaluation')}. An instrumented test runs the application's own classes on the phone: the detector factory with the shipped model files and settings, and the face path (detection, quality gate, levelled crop, embedding). "
           "Images are read from storage instead of the camera, so that every configuration sees exactly the same input. Raw detections, embeddings and per-image times are written on the phone and analysed on a laptop. "
           "The benchmark has to stay in the foreground: the phone's power manager froze a background benchmark process within seconds.")
    d.figure(paths["evaluation"], "evaluation", "Evaluation procedure: public datasets are processed on the phone by the application's own pipeline, and the raw output is analysed on a laptop.")
    d.para("For detection, each of the four models was run on the GPU and on the CPU where supported (the int8 model runs on the CPU only), seven configurations in total. "
           "We report COCO AP@[.5:.95] and AP50 at the deployed settings (detector score 0.3, at most 10 results), and precision and recall at the default announcement thresholds (0.7, people 0.5) with greedy matching at an intersection over union of 0.5; detections on crowd regions are ignored. "
           "Recall is also given by COCO object size (small below 32², medium below 96², large from 96² pixels). Latency is the duration of the detector call, without the first 20 images of each run.")
    d.para(f"For faces, the labelled person is taken to be the face nearest the image centre. Verification uses the standard pairs; a face rejected by the quality gate counts as a non-match. "
           f"Open-set identification follows the way the application is used: <i>G</i> saved people are drawn at random from the {faces['identification']['people_eligible']} identities with at least eight usable images, five images of each are enrolled, "
           "up to five further images of each are genuine probes, and the same number of images of people who are not saved are stranger probes; 300 random trials are run for each setting. The matching rule of Algorithm 2 was re-implemented in the analysis script.")
    d.para("All intervals are 95%. Precision, recall and identification rates use a bootstrap over images or trials with 2,000 resamples;[[efron1979]] verification rates use the Wilson interval.[[wilson1927]]")

    d.heading("Object detection", level=2)
    d.caption(T("detection"), f"Object detection on the phone, {L2['images']:,} COCO val2017 images. Precision and recall at the application's thresholds with 95% confidence intervals. Confidence is the mean detector score of true positives (TP) and false positives (FP).")
    rows = []
    for key in ORDER:
        e = det.get(key)
        if e:
            rows.append([NAMES[e["model"]], e["delegate"], f"{100 * e['ap']:.1f}", f"{100 * e['ap50']:.1f}", pct(e["precision"], e["precision_ci"]), pct(e["recall"], e["recall_ci"]),
                         pct(e["recall_by_size"]["large"]), f"{e['confidence_tp']:.2f} / {e['confidence_fp']:.2f}", f"{e['ms_mean']:.0f} / {e['ms_p95']:.0f}", f"{1000 / e['ms_mean']:.1f}"])
    d.table(["Model", "Delegate", "AP", "AP50", "Precision", "Recall", "Large-object recall", "Confidence TP / FP", "Latency ms, mean / p95", "FPS"], rows, [2.2, 1.5, 0.9, 1.1, 2.2, 2.2, 1.4, 1.9, 1.6, 1.0], size=8)
    d.para(f"{T('detection')} shows that precision lies between {pct(min(precisions))} and {pct(max(precisions))} for every configuration: at the announcement thresholds the application rarely names something that is not there. "
           f"Recall is low by design. The default model, EfficientDet-Lite2 on the GPU, finds {pct(L2['recall'])} of all annotated objects and {pct(L2['recall_by_size']['large'])} of the large ones ({F('sizes')}); small objects are almost never reported. "
           f"True positives have a mean detector score of {L2['confidence_tp']:.2f} and false positives of {L2['confidence_fp']:.2f}.")
    d.para(f"The default model needs {L2['ms_mean']:.0f} ms per image ({1000 / L2['ms_mean']:.1f} frames per second), and the GPU makes it {L2C['ms_mean'] / L2['ms_mean']:.1f} times faster than the CPU without changing its accuracy ({F('latency')}, {F('tradeoff')}). "
           f"The int8 model used in walk mode needs {I8['ms_mean']:.0f} ms ({1000 / I8['ms_mean']:.1f} frames per second) on the CPU at {100 * I8['ap']:.1f} AP, against {100 * L0['ap']:.1f} AP for its floating-point version.")
    d.figure(paths["latency"], "latency", "Mean inference time per image for each model on the GPU and the CPU. The int8 model runs on the CPU only.")
    d.figure(paths["tradeoff"], "tradeoff", "Accuracy against speed for the seven configurations (SSD = SSD MobileNet V2). The delegate changes the speed but not the accuracy.")
    d.figure(paths["sizes"], "sizes", "Recall at the application's thresholds by COCO object size, on the delegate each model uses by default.")
    d.para(f"The announcement threshold is a setting of the application, and {F('threshold')} shows what it trades. At 0.3 the default model reaches a recall of {pct(sweep[0]['recall'])} at a precision of {pct(sweep[0]['precision'])}; "
           f"each step of 0.1 gives up recall for precision, down to {pct(sweep[-1]['recall'])} and {pct(sweep[-1]['precision'])} at the default of 0.7. "
           f"{F('classes')} breaks the default setting down by class. People, by far the most frequent class, are found in {pct(person['recall'])} of cases at a precision of {pct(person['precision'])}; "
           f"for furniture and small household objects recall is much lower, for example {pct(chair['recall'])} for chairs, because most of these objects in COCO are small, distant or partly hidden.")
    d.figure(paths["threshold"], "threshold", "Precision and recall of the default model (EfficientDet-Lite2, GPU) for the announcement thresholds the application allows.")
    d.figure(paths["classes"], "classes", "Precision and recall of the default model at the deployed thresholds for the 15 most frequent classes of the sample (n = annotated objects). For classes with a recall of a few per cent, precision rests on fewer than ten announcements.")

    d.heading("Face recognition", level=2)
    d.caption(T("facepipe"), f"Face pipeline on the phone, LFW ({faces['images']:,} images).")
    d.table(["Measure", "Value"], [
        ["Face found", pct(faces["face_found"])],
        ["Face passes the quality gate and is embedded", pct(faces["embedded"])],
        ["Verification accuracy at the deployed threshold of 0.5 (6,000 pairs)", pct(ver["accuracy_at_0.5"])],
        ["True accept rate at 0.5", pct(ver["true_accept"], ver["true_accept_ci"])],
        ["False accept rate at 0.5", pct(ver["false_accept"], ver["false_accept_ci"])],
        ["Verification accuracy, standard 10-fold protocol", f"{100 * ver['ten_fold_accuracy']:.2f}% ± {100 * ver['ten_fold_sd']:.2f}"],
        ["Time per face, median: detection + embedding", f"{faces['find_ms_median']:.0f} + {faces['embed_ms_median']:.0f} ms"],
    ], [11.0, 5.0], size=9)
    d.para(f"A face was found in every LFW image and {pct(faces['embedded'])} of the faces passed the quality gate ({T('facepipe')}). Under the standard 10-fold protocol the verification accuracy is {100 * ver['ten_fold_accuracy']:.2f}%. "
           f"At the deployed threshold the pipeline accepts {pct(ver['true_accept'])} of genuine pairs and {pct(ver['false_accept'])} of impostor pairs. {F('scores')} shows why the threshold separates the two groups: "
           f"the median similarity is {scores['same_median']:.2f} for pairs of the same person and {scores['different_median']:.2f} for pairs of different people. One face takes {faces['find_ms_median']:.0f} ms to find and {faces['embed_ms_median']:.0f} ms to embed.")
    d.figure(paths["scores"], "scores", f"Cosine similarity of the LFW verification pairs measured on the phone ({sum(scores['same']):,} pairs of the same person and {sum(scores['different']):,} of different people; {scores['pairs_without_embedding']} pairs with a face rejected by the quality gate are not shown).")
    d.para(f"{T('ablation')} and {F('ablation')} isolate the contribution of each part of the matching rule with ten saved people. A rule that accepts the best single sample at a cosine of 0.3, the rule the project started from, names a saved person correctly in {pct(rule[0]['named_right'])} of cases but also gives a saved name to {pct(rule[0]['stranger_named'])} of strangers. "
           f"Raising the threshold to 0.5 lowers that share to {pct(rule[1]['stranger_named'])}, scoring a person by the mean of the three best samples to {pct(rule[2]['stranger_named'])}, and the margin over the runner-up to {pct(rule[3]['stranger_named'])}, "
           f"while correct naming falls by only {100 * (rule[0]['named_right'] - rule[3]['named_right']):.1f} percentage points.")
    d.caption(T("ablation"), "Ablation of the face matching rule: 10 saved people, 5 enrolled images each, 300 random trials per row, 95% confidence intervals in brackets.")
    d.table(["Matching rule", "Saved person named correctly", "Given another saved person's name", "Stranger given a saved name"],
            [[r["rule"], pct(r["named_right"], r["named_right_ci"]), pct(r["named_wrong"], r["named_wrong_ci"]), pct(r["stranger_named"], r["stranger_named_ci"])] for r in rule],
            [6.4, 3.4, 3.0, 3.2], size=9, first_bold=True)
    d.figure(paths["ablation"], "ablation", "Effect of each part of the matching rule; whiskers are 95% confidence intervals. The left axis does not start at zero.")
    d.para(f"The number of saved people matters ({T('gallery')}, {F('gallery')}). Saved people are named correctly in {pct(ident['20']['named_right'])} to {pct(ident['1']['named_right'])} of cases and are practically never confused with each other, "
           f"but every additional saved person is one more chance of a false match: the share of strangers who receive a saved name rises from {pct(ident['1']['stranger_named'])} with one saved person to {pct(ident['10']['stranger_named'])} with ten and {pct(ident['20']['stranger_named'])} with twenty.")
    d.caption(T("gallery"), "Open-set identification with the deployed rule by number of saved people: 5 enrolled images each, 300 random trials per row, 95% confidence intervals in brackets.")
    d.table(["Saved people", "Saved person named correctly", "Given another saved person's name", "Stranger given a saved name", "Genuine / stranger probes"],
            [[g, pct(r["named_right"], r["named_right_ci"]), pct(r["named_wrong"], r["named_wrong_ci"], 2), pct(r["stranger_named"], r["stranger_named_ci"]), f"{r['genuine_probes']:,} / {r['stranger_probes']:,}"] for g, r in ident.items()],
            [2.2, 3.8, 3.6, 3.6, 2.8], size=9)
    d.figure(paths["gallery"], "gallery", "Open-set identification against the number of saved people; whiskers are 95% confidence intervals. The two panels have different vertical scales, and the left one does not start at zero.")
    d.para(f"{T('tuning')} and {F('facethreshold')} show two further settings. A threshold of 0.6 instead of 0.5 would lower the share of strangers named to {pct(tune[2]['stranger_named'])} while still naming {pct(tune[2]['named_right'])} of saved people correctly; at 0.7 correct naming drops to {pct(tune[3]['named_right'])}. "
           f"More enrolled images raise correct naming, from {pct(enrol[0]['named_right'])} with one image to {pct(enrol[3]['named_right'])} with ten. The stranger rates of these rows are not directly comparable, because fewer identities have enough images for the larger enrolments.")
    d.caption(T("tuning"), "Threshold of the deployed rule and number of enrolled images, 10 saved people, 95% confidence intervals in brackets.")
    d.table(["Setting", "Saved person named correctly", "Stranger given a saved name", "Eligible identities"],
            [[f"Threshold {r['threshold']}, 5 images", pct(r["named_right"], r["named_right_ci"]), pct(r["stranger_named"], r["stranger_named_ci"]), r["people_eligible"]] for r in tune] +
            [[f"{r['enrolled']} enrolled image{'s' if r['enrolled'] > 1 else ''}, threshold 0.5", pct(r["named_right"], r["named_right_ci"]), pct(r["stranger_named"], r["stranger_named_ci"]), r["people_eligible"]] for r in enrol],
            [5.2, 4.0, 4.0, 2.8], size=9, first_bold=True)
    d.figure(paths["facethreshold"], "facethreshold", "Trade-off of the threshold on the person score with ten saved people.")

    d.heading("Comparison with published figures", level=2)
    d.para(f"{T('published')} sets the measurements beside the figures published for the same models.[[tfblog,mp_objdet]] They are not directly comparable: the published accuracy is for the whole of COCO val2017 without the deployed limits on score and number of results, and the published latency is for a Google Pixel 6. "
           f"The measured AP is {33.97 - 100 * L2['ap']:.1f} and {25.69 - 100 * L0['ap']:.1f} points below the published value for EfficientDet-Lite2 and Lite0, as the deployed limits lead one to expect, "
           f"and on the GPU the mid-range phone is {SSD['ms_mean'] / 24.01:.1f} to {L2['ms_mean'] / 41.15:.1f} times slower than the Pixel 6. "
           f"For faces, the measured verification accuracy of {100 * ver['ten_fold_accuracy']:.2f}% compares with 98.4% reported for FaceNet-512 by an independent benchmark[[deepface]] and 99.63% in the original FaceNet paper;[[facenet]] the face detector, the crop and the quality gate differ, and we did not investigate the difference further.")
    d.caption(T("published"), "This work against published figures for the same models.")
    rows = []
    for model, (published_ap, published_ms) in PUBLISHED.items():
        gpu, cpu = det.get((model, "GPU")), det.get((model, "CPU"))
        rows.append([NAMES[model], " / ".join(f"{100 * e['ap']:.1f}" if e else "n/a" for e in (gpu, cpu)), published_ap, " / ".join(f"{e['ms_mean']:.0f}" if e else "n/a" for e in (gpu, cpu)), published_ms])
    d.table(["Model", "AP, this work (GPU / CPU)", "Published COCO mAP[[tfblog]]", "Latency ms, this work (GPU / CPU)", "Published latency ms, Pixel 6 (GPU / CPU)[[mp_objdet]]"], rows, [3.4, 3.0, 3.0, 3.2, 3.4], size=9, first_bold=True)

    # ---------------- 5 Implementation ----------------
    d.heading("Environment setup and implementation")
    d.para(f"Nungil is written in Kotlin {v['kotlin']} for Android {('7.0' if v['min'] == '24' else 'API ' + v['min'])} and later, with CameraX {v['camerax']}, MediaPipe Tasks, ML Kit, TensorFlow Lite, ARCore {v['arcore']} and Room {v['room']}. "
           "It consists of 165 Kotlin source files (16,046 lines). The rule layer of 81 files (6,242 lines) has no Android dependencies and is covered by 733 unit tests, which all pass. "
           "Model files are downloaded at build time. The only function that uses the network is an optional request for a walking route, which falls back to an offline direction beacon; face embeddings and saved objects are stored in a local database on the phone.")
    d.para("The interface is built for use without sight: one large button and voice control reach every function, everything spoken is also shown as a caption, the screens have light, dark and high-contrast themes, and every control carries a spoken label for the screen reader. "
           f"{F('screens')} shows the main screens.")
    d.para("A typical session runs as follows. The user says \"Eye, look around\" and turns once on the spot; the coverage ring fills, and the phone says \"Around you: 3 blue chairs in front; a black laptop on your right.\" "
           "\"Find my bag\" starts the beeps, which speed up as the camera points at the bag. \"Who is this\" names a saved person, and \"walk mode\" switches to depth-based warnings such as \"Wall ahead, 3 steps\".")
    paragraph = d.doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rich(paragraph, "{{Screenshots of the home screen, a full scan, find mode, a saved person and walk mode}}", count=False)
    d.caption(F("screens"), "Main screens of Nungil. {{to be added}}")
    d.para("The benchmark is part of the repository: a script downloads the datasets, an instrumented test runs on the phone, and two scripts compute every number, table and figure of this paper from the raw output.")

    # ---------------- 6 Discussion ----------------
    d.heading("Discussion")
    d.para(f"The results describe an assistant that is trustworthy but conservative. At its default setting about one announcement in twenty is wrong, and what it does not see it does not invent; in exchange it reports roughly half of the large objects and very few small ones. "
           f"For a user who turns slowly through a room this is a reasonable balance, and the setting can be lowered to 0.3 for a recall of {pct(sweep[0]['recall'])} at a precision of {pct(sweep[0]['precision'])}. "
           f"At {1000 / L2['ms_mean']:.1f} frames per second the default model suits a deliberate scan but not walking, which is why walk mode uses the int8 model at {1000 / I8['ms_mean']:.1f} frames per second together with depth.")
    d.para(f"For faces the experiments expose an effect that a verification accuracy hides. A false accept rate of {pct(ver['false_accept'])} per pair looks negligible, but with ten saved people {pct(ident['10']['stranger_named'])} of strangers receive a saved name, and the share grows with every person added. "
           f"The matching rule already removes most of the error of a single-sample rule, and the measurements suggest a simple further step: a threshold of 0.6 would reduce the share to {pct(tune[2]['stranger_named'])} at a cost of about three percentage points in correct naming.")
    d.para("Compared with the literature, Nungil needs the least hardware and offers the widest set of functions on the device, while several published systems are ahead in two respects: they were tested with visually impaired participants, and those built on vision-language models can answer open questions about a scene, which a detector with 80 classes cannot. "
           "Because no shared benchmark is used in the field, we could not rank accuracy; we therefore report on public datasets at the deployed settings, so that later systems can be compared with these numbers.")
    d.heading("Limitations", level=2)
    d.bullets([
        "One phone and one run per configuration; run-to-run variation and other devices were not measured.",
        "Still images from public datasets, not live camera frames: motion blur, focus, lighting and the multi-frame confirmation of the scan are not covered.",
        "The identification experiment replays the matching rule in the analysis script with five LFW images per person; enrolment in the application collects different samples.",
        "Colour naming, find-object guidance, saved objects, the reader and walk mode are implemented but were not evaluated.",
        "The comparison with other systems and applications rests on what their papers and product pages state, not on our own measurements of them.",
        "No study with blind or low-vision users has been carried out.",
        "The application does not read Korean text, recognises a fixed vocabulary of 80 classes, runs on Android only and needs a phone with ARCore depth for walk mode.",
    ])

    # ---------------- 7 Conclusion ----------------
    d.heading("Conclusion")
    d.para(f"This paper presented Nungil, a smartphone assistant for blind and low-vision people that performs all recognition on the device, and measured its deployed pipeline on public benchmarks. "
           f"On a mid-range phone the default detector is right in {pct(L2['precision'])} of its announcements and finds {pct(L2['recall_by_size']['large'])} of large objects at {L2['ms_mean']:.0f} ms per image; "
           f"saved people are named correctly in {pct(ident['10']['named_right'])} of cases with ten people saved, and a three-part matching rule reduces the share of strangers who are given a saved name from {pct(rule[0]['stranger_named'])} to {pct(rule[3]['stranger_named'])}. "
           "The benchmark, the raw results and the analysis are open, so that the numbers can be reproduced and other systems can be measured in the same way.")
    d.para("Future work will evaluate the functions that were not measured here, in particular colour naming, find-object guidance and the depth-based walking hazards; add Korean text recognition; raise recall for small objects; and, most importantly, study the application with blind and low-vision users.")

    # ---------------- Statements ----------------
    d.heading("Statements and declarations")
    for title, text in (
        ("Ethics approval", "This work did not involve human participants or animals. All experiments used publicly available image datasets."),
        ("Consent to participate", "Not applicable."),
        ("Funding", "{{Funding statement, or: The authors received no financial support for the research, authorship, and/or publication of this article.}}"),
        ("Declaration of conflicting interests", "{{To be confirmed by all authors: The authors declared no potential conflicts of interest with respect to the research, authorship, and/or publication of this article.}}"),
        ("Data availability", "The datasets used in this work are publicly available and are cited in the manuscript. The scripts that produce every table and figure from the benchmark output are available in the code repository. {{The raw benchmark output will be attached to the tagged release.}}"),
        ("Code availability", "The source code of Nungil and of the benchmark is available at https://github.com/androidIkrom/hackathon-ml-scanner under the {{Apache-2.0}} licence."),
        ("Author contributions", "{{CRediT roles for each author: conceptualisation, software, methodology, investigation, formal analysis, writing, supervision.}}"),
        ("Use of generative AI", "{{To be confirmed by the authors: During the preparation of this work the authors used Claude (Anthropic) to write the benchmark and analysis scripts, to prepare the figures and to draft the text. The authors reviewed and edited the content and take full responsibility for the content of the publication.}}"),
    ):
        d.para(f"<b>{title}.</b> {text}", count=False)

    d.heading("Code metadata")
    d.caption(T("metadata"), "Code metadata.")
    d.table(["Nr", "Code metadata description", "Metadata"], [
        ["C1", "Current code version", "{{v1.0: tag to be created}}"],
        ["C2", "Permanent link to code/repository used for this code version", "https://github.com/androidIkrom/hackathon-ml-scanner"],
        ["C3", "Legal code license", "{{Apache-2.0: licence file to be added to the repository}}"],
        ["C4", "Code versioning system used", "git"],
        ["C5", "Software code languages, tools and services used", "Kotlin, Android SDK, CameraX, MediaPipe Tasks, ML Kit, TensorFlow Lite, ARCore, Room; Python (benchmark analysis); optional openrouteservice routing"],
        ["C6", "Compilation requirements, operating environments and dependencies", f"JDK 17 or later, Android SDK {v['target']}, Gradle with Android Gradle Plugin {v['agp']}; Android {('7.0' if v['min'] == '24' else 'API ' + v['min'])} or later with a gyroscope and a compass; ARCore depth for walk mode"],
        ["C7", "If available, link to developer documentation/manual", "https://github.com/androidIkrom/hackathon-ml-scanner (README and docs/build-guide.md)"],
        ["C8", "Support email for questions", "{{support email}}"],
    ], [1.2, 6.3, 8.5], size=9)

    body_words = WORDS[0]
    d.heading("References")
    for number, key in enumerate(CITES.order, 1):
        paragraph = d.doc.add_paragraph()
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.left_indent, paragraph.paragraph_format.first_line_indent = Cm(0.9), Cm(-0.9)
        rich(paragraph, f"{number}. {REFS[key]}", size=10, count=False)
    unused = [key for key in REFS if key not in CITES.order]

    # ---------------- Notes for the authors ----------------
    d.page_break()
    d.heading("Notes for the authors (remove before submission)")
    d.para("<b>Highlights</b> (3 to 5 bullet points of at most 85 characters, required by Elsevier journals):", count=False)
    d.bullets(highlights, count=False)
    d.para("<b>Blanks to fill</b> (highlighted in yellow in the text):", count=False)
    d.bullets(["Authors, affiliations and corresponding author.", "Funding, conflict-of-interest statement and author contributions.", "Licence: add the licence file to the repository and create the version tag named in the code metadata.",
               "Data availability: publish the raw benchmark output (about 30 MB) with the tagged release, or reword the statement.",
               f"{F('screens')}: screenshots of the application.", "Support email in the code metadata.", "Statement on the use of generative AI: confirm or reword."], count=False)
    d.para("<b>To verify before submission:</b>", count=False)
    d.bullets([
        "The cells of the related-work table were extracted from the papers by an automated literature search and quoted from their text; the bibliographic data of every reference were checked against Crossref, PubMed or arXiv, but the reported numbers should be checked once more against the PDFs.",
        "On 12 October 2026 every quoted number of ten of the twelve systems was confirmed against the full text (Europe PMC or arXiv); See et al. 2022 and DeepNAVI could not be retrieved (HTTP 403) and still rest on the original extraction. Bai et al. 2019, Chen et al. 2024 and WalkVLM were read as arXiv versions.",
        "Two product pages (Be My AI help, Sullivan+ on Google Play) could not be opened directly and rest on search-result excerpts; several cells of the applications table are marked unverified.",
        "The GPU name (Mali-G57) and the amount of memory (8 GB) of the test phone come from the project documentation, not from a measurement on the phone.",
        "The statement that a cosine of 0.3 on one sample was the project's initial rule comes from the project's source comments.",
        "The counts of source files, lines and unit tests are those of commit b65c9b3 plus the local changes present on 3 October 2026; recount them for the tagged version.",
        "Journal requirements below come from search summaries of the author guidelines, whose pages could not be opened directly; check them on the journal sites.",
    ], count=False)
    d.para("<b>Fit to the two candidate journals:</b>", count=False)
    d.bullets([
        f"Technology and Health Care (Sage): structured abstract of about 200 words (the abstract here has {abstract_words(d)} words), 4 to 7 keywords, Sage Vancouver references (used here), disclosure of AI tools in the methods or acknowledgements.",
        f"SoftwareX (Elsevier): original software publications are limited to about 4,000 words and 6 figures, with the sections Motivation and significance, Software description, Illustrative examples, Impact and Conclusions; this manuscript has about {body_words:,} words and {len(FIGURES)} figures and would have to be shortened. The code metadata table, highlights and the declaration of generative AI are already included.",
    ], count=False)
    if unused:
        d.para("References defined but not cited: " + ", ".join(unused), count=False)
    d.doc.save(target)
    return body_words, len(CITES.order), unused


def abstract_words(d):
    count, inside = 0, False
    for paragraph in d.doc.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            inside = paragraph.text == "Abstract"
            continue
        if inside and not paragraph.text.startswith("Keywords"):
            count += len(paragraph.text.split())
    return count


if __name__ == "__main__":
    data = Path(sys.argv[1] if len(sys.argv) > 1 else "D:/nungil-bench")
    target = data / (sys.argv[2] if len(sys.argv) > 2 else "Nungil-manuscript.docx")
    results = json.loads((data / "results.json").read_text())
    words, references, unused = build(results, figures(results, data / "figures"), target)
    print("wrote", target, "| words in the body:", words, "| references:", references, "| uncited:", unused)
